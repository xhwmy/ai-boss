"""
简历附件上传模块（v3.2）
支持 .docx / .pdf 简历解析，提取的全文会在 AI 评分时合并补充到结构化简历文本中。

注意：本模块不 import config，避免循环依赖（config.build_resume_text 会调用本模块）。
"""

import json
import time
from pathlib import Path

# 项目根目录
ROOT_DIR = Path(__file__).parent

# 存储位置
UPLOAD_DIR = ROOT_DIR / "uploads"
DATA_DIR = ROOT_DIR / "data"
META_FILE = DATA_DIR / "uploaded_resume.json"
TEXT_FILE = DATA_DIR / "uploaded_resume.txt"

# 限制
ALLOWED_EXTS = {".docx", ".pdf"}
MAX_FILE_MB = 10
MAX_FILE_BYTES = MAX_FILE_MB * 1024 * 1024
MAX_TEXT_CHARS = 8000


def parse_docx(path) -> str:
    """解析 .docx：段落 + 表格"""
    from docx import Document
    doc = Document(str(path))
    parts = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            parts.append(text)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            line = " | ".join(x for x in cells if x)
            if line:
                parts.append(line)
    return "\n".join(parts)


def parse_pdf(path) -> str:
    """解析 .pdf：逐页提取文本"""
    import pdfplumber
    parts = []
    with pdfplumber.open(str(path)) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            if text.strip():
                parts.append(text.strip())
    if not parts:
        raise ValueError("PDF 中未提取到文字（可能是图片型扫描件），请改用文字版 PDF 或 .docx")
    return "\n".join(parts)


def parse_resume_file(path, ext: str) -> str:
    """按扩展名解析简历文件为纯文本，不支持的格式抛出友好 ValueError"""
    ext = (ext or "").lower()
    if ext == ".docx":
        return parse_docx(path)
    if ext == ".pdf":
        return parse_pdf(path)
    raise ValueError("仅支持 .docx 和 .pdf 格式，请将文件另存为 .docx 或 PDF")


def _truncate(text: str, max_chars: int = MAX_TEXT_CHARS) -> str:
    """截断过长的简历文本，避免超出 token 上限"""
    if len(text) <= max_chars:
        return text
    return text[:max_chars] + f"\n……（简历过长，已截断至前 {max_chars} 字）"


def save_uploaded(file) -> dict:
    """保存上传的简历附件并提取文本。

    参数 file: werkzeug.datastructures.FileStorage
    返回 meta 字典；校验或解析失败抛 ValueError。
    """
    from werkzeug.utils import secure_filename

    original_name = file.filename or ""
    ext = Path(original_name).suffix.lower()
    if ext not in ALLOWED_EXTS:
        raise ValueError("仅支持 .docx 和 .pdf 格式，请将文件另存为 .docx 或 PDF")

    # 读取文件内容并校验大小
    data = file.read()
    if not data:
        raise ValueError("文件内容为空")
    if len(data) > MAX_FILE_BYTES:
        raise ValueError(f"文件过大（>{MAX_FILE_MB}MB），请压缩后重试")

    # 落盘文件名用 secure_filename 防路径穿越；中文名会被剥离，改用安全兜底名
    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    safe = secure_filename(original_name)
    if not safe or Path(safe).suffix.lower() != ext:
        safe = f"resume_{int(time.time())}{ext}"
    saved_name = safe
    saved_path = UPLOAD_DIR / saved_name
    if saved_path.exists():
        saved_name = f"{Path(saved_name).stem}_{int(time.time())}{ext}"
        saved_path = UPLOAD_DIR / saved_name
    saved_path.write_bytes(data)

    # 提取文本；解析失败则删除已落盘文件，不留残余
    try:
        raw_text = parse_resume_file(saved_path, ext)
    except Exception:
        saved_path.unlink(missing_ok=True)
        raise

    was_truncated = len(raw_text) > MAX_TEXT_CHARS
    text = _truncate(raw_text)

    # 写入提取文本 + 元数据
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    TEXT_FILE.write_text(text, encoding="utf-8")
    meta = {
        "filename": original_name,
        "saved_name": saved_name,
        "format": ext.lstrip("."),
        "uploaded_at": time.strftime("%Y-%m-%d %H:%M:%S"),
        "size_bytes": len(data),
        "text_length": len(text),
        "truncated": was_truncated,
    }
    META_FILE.write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")
    return meta


def get_uploaded_meta() -> dict | None:
    """读取上传简历的元数据；无上传返回 None"""
    if not META_FILE.exists():
        return None
    try:
        with open(META_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return None


def get_uploaded_text() -> str | None:
    """读取上传简历的提取文本；无上传返回 None"""
    if not TEXT_FILE.exists():
        return None
    try:
        return TEXT_FILE.read_text(encoding="utf-8")
    except Exception:
        return None


def clear_uploaded() -> None:
    """清空上传的简历附件（删文件 + 删元数据 + 删提取文本）"""
    meta = get_uploaded_meta()
    if meta and meta.get("saved_name"):
        try:
            (UPLOAD_DIR / meta["saved_name"]).unlink(missing_ok=True)
        except Exception:
            pass
    META_FILE.unlink(missing_ok=True)
    TEXT_FILE.unlink(missing_ok=True)
